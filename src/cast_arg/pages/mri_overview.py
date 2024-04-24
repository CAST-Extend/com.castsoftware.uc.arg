from cast_arg.pages.mri_report import MRIPage
from cast_arg.powerpoint import PowerPoint
from cast_common.util import format_table,list_to_text
from copy import deepcopy

from pandas import Series
from pptx.dml.color import RGBColor

class MRIOverview(MRIPage):
    description = 'MRI Overview Page'

    def report(self,config) -> bool:
        app_cnt = len(config.application)
        table = self.ppt.get_shape_by_name('Table 3')
        if table is None:
            raise ValueError(f'Table not found in template: ')
        table = table.table
        new_row = deepcopy(table._tbl.tr_lst[-1])                     

        for idx in range(1,app_cnt):
            last_row = len(table.rows)-1
            table._tbl.append(new_row)
            new_row = last_row + 1

            for col in range(0,len(table.columns)):
                new_cell = table.cell(new_row,col) 
                for p in new_cell.text_frame.paragraphs:
                    self.ppt._replace_paragraph_text(p, '{app1',f'{{app{idx+1}')
                pass
            pass